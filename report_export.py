"""
Maps Account_research_automation.py's internal result columns onto the
column structure of the reference file Faheem attached
(Voiant_Prospect_List_targets_MASTER_backup_20260724.xlsx), for the app's
Faheem-facing download.

IMPORTANT -- THIS IS DISPLAY-ONLY, NOT A SCORING CHANGE
---------------------------------------------------------
Nothing here recomputes Priority Score/Band, re-judges Platform
Confidence, or touches the Sales-vs-Finance classification. Those numbers
come straight out of Account_research_automation.py's own
compute_priority_score() / apply_parsed_result(), unchanged. This module
only renames/combines columns for presentation.

A REAL CAVEAT THIS FILE CANNOT PAPER OVER: the reference file's Priority
Score is out of 100 (Alex's own weighted formula: Switching Signals 35% /
Company Fit 25% / Timing 20% / Engagement 20%). This build's Priority
Score is out of ~10 (Section 7 of the brief: Trigger 30% / Platform
Confidence 25% / Size 20% / Industry 15% / Recency 10%). Matching the
COLUMN NAME does not make the two SCALES comparable -- a "6" here and a
"60" there are not the same claim. The Read Me sheet below says this
explicitly so Faheem doesn't read them side by side as equivalent.
"""

import pandas as pd

import xlsx_style

REFERENCE_COLUMNS = [
    "#",
    "Company",
    "Industry",
    "HQ",
    "Est. Employees",
    "Anaplan Use Case\n(Sales Planning Focus)",
    "Key Decision-Maker\n(Role to Target)",
    "Switching Signal / Trigger",
    "Priority\nScore",
    "Priority\nBand",
    "Data Reliability\n(Use Case + Trigger)",
    "Evidence Source",
    "Status",
    "Website",
    "Notes / Next Steps",
]

READ_ME_TEXT = [
    ["This sheet's columns are formatted to match Alex's reference file for easy side-by-side review."],
    [""],
    ["Priority Score / Band are NOT on the same scale as the reference file:"],
    ["  - This file: Section 7 brief formula, roughly 0-10 (Trigger 30% / Platform Confidence 25% / Size 20% / Industry 15% / Recency 10%)"],
    ["  - Reference file: Alex's formula, 0-100 (Switching Signals 35% / Company Fit 25% / Timing 20% / Engagement 20%)"],
    ["  Same column name, different math -- do not compare the numbers directly across the two files."],
    [""],
    ["'Anaplan Use Case' column repurposed here to show Platform + Department findings (this brief covers Anaplan/Pigment/Varicent, not an Anaplan-only migration play)."],
    ["'Website' is not currently captured by this pipeline -- left blank (known limitation, not a bug)."],
    ["Est. Employees is blank unless manually added to the input list -- scoring defaults to the lowest size bucket until then."],
]

# Same Rating/Label/What It Means/How To Use structure as Alex's reference
# file's own Reliability Key sheet, but describing OUR criteria (Section 4's
# Platform Confidence rubric + whether anything was flagged shaky during
# research), not her Anaplan.com-source-specific wording -- the two files'
# reliability tiers are not measuring the same thing even where the labels
# happen to match.
RELIABILITY_KEY_ROWS = [
    ["Data Reliability Scoring -- How To Read This List"],
    [""],
    ["Rating", "Label", "What It Means", "How To Use In Outreach"],
    ["★★★", "Verified",
     "Platform Confidence = Confirmed, and nothing about the finding was flagged as shaky. Both WHICH tool is used and that it's on the sales/RevOps side (not Finance-only) are backed by a public source (case study, job posting, press, or conference mention).",
     "Highest confidence. Safe to reference the specific finding in outreach (the job posting, the case study, etc)."],
    ["★★", "Corroborated",
     "Platform Confidence = Confirmed, but the research flagged something shaky along the way -- e.g. the named contact may be stale/departed, or one piece of supporting evidence was ambiguous. See the Notes column for the specific caveat.",
     "Good confidence, but verify the flagged caveat first. Don't lead with a detail that might be stale."],
    ["★", "Partial",
     "Platform Confidence = Inferred. The tool is named in a public source, but whether it's used on the sales side (vs Finance/FP&A-only) isn't clearly confirmed.",
     "Use cautiously. Don't claim sales-side usage in first outreach -- lead with the business problem and confirm the department during discovery."],
    ["◇", "Inferred / Hypothesis",
     "Platform Confidence = Unconfirmed. No credible public mention of Anaplan, Pigment, or Varicent tied to this company at all. Included because it fits the ICP profile (industry, size, geography), not because usage is confirmed.",
     "Do NOT reference any specific platform in outreach -- you don't know if they use one. Lead entirely with the business pain point; confirm the actual tool stack during discovery."],
]


def _s(value) -> str:
    """NaN-safe stringify -- pandas hands back float('nan') for blank Excel
    cells, which breaks str.join() and string formatting outright."""
    if value is None:
        return ""
    if isinstance(value, float) and pd.isna(value):
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def _data_reliability(row) -> str:
    confidence = _s(row.get("Platform Confidence (Confirmed/Inferred/Unconfirmed)"))
    has_flag = bool(_s(row.get("Low Confidence Reason")))

    if confidence == "Confirmed":
        return "Corroborated" if has_flag else "Verified"
    if confidence == "Inferred":
        return "Partial"
    return "Inferred/Hypothesis"


def _status(row) -> str:
    if _s(row.get("Error")):
        return "Error - Needs Manual Research"
    if row.get("Outreach Ready"):
        return "Outreach Ready"
    if row.get("NEEDS_MANUAL_REVIEW"):
        return "Needs Review"
    return "Researched"


def _use_case_summary(row) -> str:
    platform = _s(row.get("Platform (Anaplan/Pigment/Varicent)")) or "None found"
    department = _s(row.get("Department Confirmed? (Sales-Ops/RevOps/Finance-only/Unknown)"))
    if department and department != "N/A":
        return f"{platform} -- {department}"
    return platform


def _decision_maker(row) -> str:
    name = _s(row.get("Named Contact"))
    title = _s(row.get("Contact Title"))
    url = _s(row.get("Contact LinkedIn URL"))
    parts = [p for p in (name, title) if p]
    text = " -- ".join(parts)
    if url:
        text = f"{text} ({url})" if text else url
    return text


def _trigger_summary(row) -> str:
    event = _s(row.get("Trigger Event"))
    source_date = _s(row.get("Trigger Source/Date"))
    if event and source_date:
        return f"{event}\nSource/Date: {source_date}"
    return event or source_date


def _notes(row) -> str:
    parts = [
        _s(row.get("Research Summary")),
        _s(row.get("Score Notes")),
        _s(row.get("Low Confidence Reason")),
        _s(row.get("Reviewer Note")),  # present on Faheem-reviewed batches, absent on raw automated output
        _s(row.get("Research Method")),  # ditto -- e.g. "Manual+Automated (agree)"
    ]
    return "\n".join(p for p in parts if p)


# --------------------------------------------------------------------------
# Faheem-reviewed batch schema (e.g. "Faheem_Review_10_Account_Batch.xlsx")
# uses simplified/renamed columns vs. the raw automated output -- this is
# the Section 10 review-gate artifact, not a different pipeline. Normalize
# it to the same canonical column names to_reference_format() expects so
# both schemas go through one mapping path.
# --------------------------------------------------------------------------

REVIEWED_SCHEMA_RENAMES = {
    "Platform": "Platform (Anaplan/Pigment/Varicent)",
    "Platform Confidence": "Platform Confidence (Confirmed/Inferred/Unconfirmed)",
    "Department": "Department Confirmed? (Sales-Ops/RevOps/Finance-only/Unknown)",
    "Trigger Strength": "Trigger Strength (Strong/Weak)",
    "Outreach Ready (Reviewed)": "Outreach Ready",
    "ICP Industry": "Industry",
}


def is_reviewed_schema(df: pd.DataFrame) -> bool:
    return "Outreach Ready (Reviewed)" in df.columns


def normalize_reviewed_schema(df: pd.DataFrame) -> pd.DataFrame:
    return df.rename(columns=REVIEWED_SCHEMA_RENAMES)


def to_reference_format_any(df: pd.DataFrame, master_context_df: pd.DataFrame = None) -> pd.DataFrame:
    """Accepts either the raw Account_Research_Results.xlsx schema or a
    Faheem-reviewed batch schema and produces the same reference-format
    output either way. The reviewed schema already carries Industry/HQ
    directly, so master_context_df is optional for it."""
    if is_reviewed_schema(df):
        df = normalize_reviewed_schema(df)
        # Reviewed files carry Industry/HQ inline -- build a same-shape
        # context lookup from the file itself instead of requiring the
        # 595-company master, unless the caller passed one explicitly.
        if master_context_df is None and "Industry" in df.columns:
            master_context_df = pd.DataFrame({
                "Company": df["Company"],
                "ICP_Bucket": df.get("Industry", ""),
                "Headquarters": df.get("HQ", ""),
                "Employees": df.get("Employees", ""),
            })
    return to_reference_format(df, master_context_df)


def to_reference_format(results_df: pd.DataFrame, master_context_df: pd.DataFrame = None) -> pd.DataFrame:
    """results_df: whatever's currently in Account_Research_Results.xlsx.
    master_context_df: optional (Company, Industry/ICP_Bucket, Headquarters/HQ,
    Employees) lookup -- the results file itself doesn't carry Industry/HQ/
    Employees, only the master input list does, so pass it in if you want
    those columns populated rather than blank."""
    context_lookup = {}
    if master_context_df is not None:
        cols = {c.strip(): c for c in master_context_df.columns}
        for _, r in master_context_df.iterrows():
            company = str(r[cols.get("Company", "Company")]).strip()
            industry = r.get(cols.get("ICP_Bucket", cols.get("Industry", "")), "")
            hq = r.get(cols.get("Headquarters", cols.get("HQ", "")), "")
            employees = r.get(cols.get("Employees", ""), "")
            context_lookup[company] = (industry, hq, employees)

    rows = []
    for i, (_, row) in enumerate(results_df.iterrows(), start=1):
        company = row.get("Company", "")
        industry, hq, employees = context_lookup.get(company, ("", "", ""))
        rows.append({
            "#": i,
            "Company": company,
            "Industry": industry,
            "HQ": hq,
            "Est. Employees": employees,
            "Anaplan Use Case\n(Sales Planning Focus)": _use_case_summary(row),
            "Key Decision-Maker\n(Role to Target)": _decision_maker(row),
            "Switching Signal / Trigger": _trigger_summary(row),
            "Priority\nScore": row.get("Priority Score", ""),
            "Priority\nBand": row.get("Priority Band", ""),
            "Data Reliability\n(Use Case + Trigger)": _data_reliability(row),
            "Evidence Source": row.get("Source Links", ""),
            "Status": _status(row),
            "Website": "",
            "Notes / Next Steps": _notes(row),
        })

    return pd.DataFrame(rows, columns=REFERENCE_COLUMNS)


def _write_three_sheet_workbook(ref_df: pd.DataFrame, out_path: str) -> str:
    """Shared by write_reference_excel() (post-research) and
    write_scanner_reference_excel() (pre-research) -- same Prospect
    List / Read Me / Reliability Key structure either way, so Faheem sees
    one consistent document shape regardless of which stage of the
    pipeline produced it."""
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        ref_df.to_excel(writer, sheet_name="Prospect List", index=False)
        pd.DataFrame(READ_ME_TEXT).to_excel(writer, sheet_name="Read Me", index=False, header=False)
        pd.DataFrame(RELIABILITY_KEY_ROWS).to_excel(writer, sheet_name="Reliability Key", index=False, header=False)
    xlsx_style.style_xlsx(out_path, sheet_names=["Prospect List"])
    xlsx_style.color_code_column(out_path, "Priority\nBand", xlsx_style.PRIORITY_BAND_COLORS, sheet_name="Prospect List")
    _style_reliability_key_sheet(out_path)
    return out_path


def write_reference_excel(results_df: pd.DataFrame, out_path: str, master_context_df: pd.DataFrame = None):
    ref_df = to_reference_format_any(results_df, master_context_df)
    return _write_three_sheet_workbook(ref_df, out_path)


def write_scanner_reference_excel(scanner_df: pd.DataFrame, out_path: str) -> str:
    """scanner_df: the priority_requeue data (Company, ICP_Bucket,
    Headquarters, Employees, Scanner Strength, Scanner Categories, Scanner
    Reason) -- this is PRE-research, so Priority Score/Band, Anaplan Use
    Case, Key Decision-Maker, and Data Reliability are genuinely blank,
    not omitted by mistake -- that data doesn't exist until Account
    Research runs. Status is 'Pending Research', the exact same
    convention the reference file itself already uses for its own
    not-yet-researched rows, so this isn't a new concept, just applied
    consistently."""
    rows = []
    strength_per_row = []
    for i, (_, row) in enumerate(scanner_df.iterrows(), start=1):
        rows.append({
            "#": i,
            "Company": row.get("Company", ""),
            "Industry": row.get("ICP_Bucket", ""),
            "HQ": row.get("Headquarters", ""),
            "Est. Employees": row.get("Employees", ""),
            "Anaplan Use Case\n(Sales Planning Focus)": "",
            "Key Decision-Maker\n(Role to Target)": "",
            "Switching Signal / Trigger": row.get("Scanner Reason", ""),
            "Priority\nScore": "",
            "Priority\nBand": "",
            "Data Reliability\n(Use Case + Trigger)": "",
            "Evidence Source": "",
            "Status": "Pending Research",
            "Website": "",
            "Notes / Next Steps": f"Scanner flagged: {row.get('Scanner Strength', '')} ({row.get('Scanner Categories', '')})",
        })
        strength_per_row.append(xlsx_style.SCANNER_STRENGTH_COLORS.get(row.get("Scanner Strength", "")))
    ref_df = pd.DataFrame(rows, columns=REFERENCE_COLUMNS)
    out_path = _write_three_sheet_workbook(ref_df, out_path)
    # Priority Band is genuinely blank here (no research yet), so
    # color_code_column has nothing to key off -- color by Scanner
    # Strength instead so the sheet isn't uncolored just because it's
    # pre-research. Same green/yellow visual language either way.
    xlsx_style.color_code_rows_by_list(out_path, strength_per_row, sheet_name="Prospect List")
    return out_path


def _style_reliability_key_sheet(path):
    """Bespoke styling, not xlsx_style.style_sheet() -- that helper assumes
    row 1 IS the header, but this sheet has a title row (1), a blank
    spacer (2), then the real header (3), so it needs its own layout."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.load_workbook(path)
    ws = wb["Reliability Key"]

    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells("A1:D1")

    header_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    for cell in ws[3]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    widths = {"A": 10, "B": 20, "C": 60, "D": 55}
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    # Same A/B/C/D palette as the Priority Band column, in order, so Faheem
    # can connect "this color in the Reliability Key" to "this color on a
    # row in Prospect List" at a glance -- Verified=A's green down to
    # Inferred/Hypothesis=D's red, same visual language across both sheets.
    row_colors = [xlsx_style.PRIORITY_BAND_COLORS["A"], xlsx_style.PRIORITY_BAND_COLORS["B"],
                  xlsx_style.PRIORITY_BAND_COLORS["C"], xlsx_style.PRIORITY_BAND_COLORS["D"]]

    for offset, row_idx in enumerate(range(4, ws.max_row + 1)):
        fill = PatternFill(start_color=row_colors[offset], end_color=row_colors[offset], fill_type="solid")
        for col_idx in range(1, 5):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.fill = fill

    ws.freeze_panes = "A4"
    wb.save(path)
